#!/usr/bin/env python3
"""Linkprüfung der Excel-Übersicht (Q-Rahmen Münchenstein / Wil / Seelisberg).

Prüft für jede Dokumentzeile (Blätter 1–6) und jeden Layer-Link (Blatt 7), ob der
Link erreichbar ist und auf das Dokument zeigt, das Bezeichnung und Referenz
nennen. Bei falschem oder unspezifischem Link wird der richtige eingetragen.

Prüfquellen:
  - Fedlex (SR-Nummer): SPARQL-Endpunkt fedlex.data.admin.ch, SR-Nummer <-> ELI.
  - LexWork-Sammlungen (SGS BL, sGS SG, RB UR, SRS Wil): JSON-API und
    Titelverzeichnis (lightweight_index) der jeweiligen Sammlung.
  - map.geo.admin.ch: Layer-ID gegen layersConfig von api3.geo.admin.ch.
  - übrige Seiten/PDFs: HTTP-Status, Titel und Text werden nach Bezeichnung durchsucht.

Ausgabe:
  linkpruefung_excel.json       Rohdaten aller Abfragen
  Linkpruefung_Excel.md         Bericht (Korrekturen, Hinweise, alle Zeilen)
  <xlsx>_geprueft.xlsx          Excel mit korrigierten Links, Link-Status und
                                neuem Blatt «9_Linkpruefung» (Option --apply)

Aufruf (läuft im GitHub-Runner, weil die Portale aus der Sandbox gesperrt sind):
  pip install openpyxl requests pypdf
  python3 linkpruefung_excel.py Q-Rahmen_Uebersicht_Muenchenstein_Wil_Seelisberg.xlsx --apply
"""
import argparse
import datetime
import io
import json
import re
import sys
import urllib.parse
from concurrent.futures import ThreadPoolExecutor
from html import unescape

import openpyxl
import requests
from openpyxl.styles import Font, PatternFill

UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) Linkpruefung Q-Rahmen",
      "Accept-Language": "de-CH,de;q=0.9"}
SPARQL = "https://fedlex.data.admin.ch/sparqlendpoint"
LAYERS_CONFIG = "https://api3.geo.admin.ch/rest/services/api/MapServer/layersConfig?lang=de"
DOC_SHEETS = ["1_Q1_Rechtsordnung", "2_Q2_Landschaft", "3_Q3_Klima_Energie",
              "4_Q4_Resilienz", "5_Q5_Ortsbild", "6_Q6_Bestellung"]
GIS_SHEET = "7_GIS_Layers"
LEXWORK = {  # Präfix der Referenz -> Basis-URL der Gesetzessammlung
    "SGS": "https://bl.clex.ch",
    "sGS": "https://www.gesetzessammlung.sg.ch",
    "RB": "https://rechtsbuch.ur.ch",
    "SRS": "https://wil-sg.tlex.ch",
}
LEXWORK_URL = re.compile(r"(https://[^/]+)/(?:app/de/texts_of_law|data)/([0-9][0-9A-Za-z.\-]*)")
STOP = set("der die das und oder von vom zum zur zu den dem des im in über für mit auf an bei "
           "nach als aus ein eine einer eines bl sg ur gde kanton kantonale kantonaler bund inkl "
           "plus sowie gesetz verordnung reglement bundesgesetz schweiz".split())

# map.geo.admin.ch: Zentrum (LV95) und Zoomstufe je Spalte in Blatt 7
MAP_VIEWS = {"F": ("2660000,1190000", 1), "G": ("2613308,1263110", 9),
             "H": ("2721138,1258009", 9), "I": ("2687333,1203124", 9)}
# Layer-IDs, die es in layersConfig nicht gibt -> richtige ID (siehe layer_suche.md)
LAYER_FIX = {
    "ch.bafu.fauna-wildtierkorridore_national": "ch.bafu.fauna-wildtierkorridor_national",
    "ch.bafu.gefahren-oberflaechenabfluss": "ch.bafu.gefaehrdungskarte-oberflaechenabfluss",
}
# Blatt 7, Zeilen ohne Layer-ID, deren Bezeichnung genau einem Layer entspricht
LAYER_ADD = {"Wasser- und Zugvogelreservate": "ch.bafu.bundesinventare-vogelreservate"}

# Manuell recherchierte Korrekturen (werden im Lauf auf Erreichbarkeit und Inhalt geprüft).
# Schlüssel: (Blatt, Zeile[, Spalte]) -> (neuer Link, Begründung, neue Referenz oder None)
KORREKTUREN = {
    ("1_Q1_Rechtsordnung", 55): (
        "https://bl.clex.ch/app/de/texts_of_law/791",
        "SGS 780 ist das Umweltschutzgesetz BL (USG BL); das DHG BL ist SGS 791", "SGS 791"),
    ("1_Q1_Rechtsordnung", 64): (
        "https://www.gesetzessammlung.sg.ch/app/de/texts_of_law/271.53",
        "sGS 271.51 ist seit 31.07.2019 aufgehoben; ersetzt durch den Nachfolgeerlass sGS 271.53",
        "sGS 271.53"),
    ("1_Q1_Rechtsordnung", 61): (
        "https://www.gesetzessammlung.sg.ch/app/de/texts_of_law/734.1",
        "Link zeigte auf die Startseite sg.ch; Wasserbaugesetz SG = sGS 734.1", "sGS 734.1"),
    ("1_Q1_Rechtsordnung", 70): (
        "https://rechtsbuch.ur.ch/app/de/texts_of_law/10.5101",
        "Link zeigte auf die Startseite ur.ch; Gesetz über den Natur- und Heimatschutz = RB 10.5101",
        "RB 10.5101"),
    ("1_Q1_Rechtsordnung", 72): (
        "https://rechtsbuch.ur.ch/app/de/texts_of_law/40.1211",
        "Link zeigte auf die Startseite ur.ch; Wasserbaugesetz Uri (WBG) = RB 40.1211", "RB 40.1211"),
    ("1_Q1_Rechtsordnung", 74): (
        "https://ur.ch/_doc/320131",
        "Link zeigte auf eine Schulungspräsentation; Gesetzestext steht in der Abstimmungsbotschaft "
        "vom 12.03.2023 (im Rechtsbuch Uri nicht als eigener Erlass geführt)", None),
    ("6_Q6_Bestellung", 23): (
        "https://ur.ch/_doc/320131",
        "Link zeigte auf eine Schulungspräsentation; Gesetzestext steht in der Abstimmungsbotschaft "
        "vom 12.03.2023 (im Rechtsbuch Uri nicht als eigener Erlass geführt)", None),
    ("1_Q1_Rechtsordnung", 88): (
        "https://wil-sg.tlex.ch/app/de/texts_of_law/7.4.3-6",
        "Link zeigte auf einen Entwurf («vom ____»); geltende Fassung = SRS 7.4.3-6", "SRS 7.4.3-6"),
    ("4_Q4_Resilienz", 26): (
        "https://wil-sg.tlex.ch/app/de/texts_of_law/7.4.3-6",
        "Link zeigte auf einen Entwurf («vom ____»); geltende Fassung = SRS 7.4.3-6", "SRS 7.4.3-6"),
    ("6_Q6_Bestellung", 11): (
        "https://www.bpuk.ch/bpuk/konkordate/ivoeb/ivoeb-2019",
        "PDF war die alte IVöB vom 15.03.2001; ersetzt durch die IVöB-2019-Seite der BPUK", None),
    ("6_Q6_Bestellung", 15): (
        "https://bl.clex.ch/app/de/texts_of_law/420",
        "Link zeigte auf die Startseite baselland.ch; EG IVöB BL = SGS 420", "SGS 420"),
    ("4_Q4_Resilienz", 17): (
        "https://www.geoportal.ch/ktsg/map/3502?y=2721138&x=1258009&scale=50000&rotation=0",
        "Kartenzentrum lag rund 6 km südlich von Wil; auf Wil zentriert", None),
    (GIS_SHEET, 35, "G"): (
        "https://geoview.bl.ch/?map_x=2613308&map_y=1263110&map_zoom=4&baselayer_opacity=100"
        "&baselayer_ref=grundkarte_farbig&tree_group_layers_Naturgefahren="
        "ghk_wassergefahren_ueberschwemmung_wildbach%2Cghk_wassergefahren_ueberschwemmung_talfluesse",
        "Koordinaten lagen bei Liestal; auf Münchenstein gesetzt", None),
    (GIS_SHEET, 41, "H"): (
        "https://www.geoportal.ch/ktsg/map/3502?y=2721138&x=1258009&scale=50000&rotation=0",
        "Kartenzentrum lag rund 6 km südlich von Wil; auf Wil zentriert", None),
    (GIS_SHEET, 43, "I"): (
        "https://geo.ur.ch/?Lat=46.9735&Layers=Nutzungsplanung%2CLiegenschaften%2CBaulanddatenbank"
        "&Lng=8.5868&Opacity=0%2C0.4%2C1&Visibility=1%2C1%2C1&Zoom=14",
        "URL war beschädigt («¢er», abgeschnittener Parameter) und Lat/Lng lag bei Altdorf; "
        "auf Seelisberg gesetzt", None),
}
# Bekannte Mängel, für die kein besserer öffentlicher Link gefunden wurde
HINWEISE = {
    ("1_Q1_Rechtsordnung", 86): "PDF ist der Entwurf zur öffentlichen Auflage (13.10.2022), nicht die "
                                "erlassene Fassung; in der Rechtssammlung Wil (tlex) nicht vorhanden",
    ("2_Q2_Landschaft", 29): "PDF ist der Entwurf zur öffentlichen Auflage (13.10.2022), nicht die "
                             "erlassene Fassung; in der Rechtssammlung Wil (tlex) nicht vorhanden",
    ("5_Q5_Ortsbild", 24): "PDF ist der Entwurf zur öffentlichen Auflage (13.10.2022), nicht die "
                           "erlassene Fassung; in der Rechtssammlung Wil (tlex) nicht vorhanden",
    ("1_Q1_Rechtsordnung", 87): "Referenz SRS 7.1-1 ist das Baureglement; eine eigene Energieplanung "
                                "ist in der Rechtssammlung Wil nicht vorhanden",
    ("3_Q3_Klima_Energie", 27): "Referenz SRS 7.1-1 ist das Baureglement; eine eigene Energieplanung "
                                "ist in der Rechtssammlung Wil nicht vorhanden",
    ("6_Q6_Bestellung", 25): "Link (RB 40.1111, PBG Uri) passt zur Bezeichnung; die Spalte Referenz "
                             "nennt dagegen IVöB/Submissionsreglement",
}
FILL_KORR = PatternFill("solid", fgColor="FFFFEB9C")
FILL_OK = PatternFill("solid", fgColor="FFC6EFCE")
FILL_WARN = PatternFill("solid", fgColor="FFFCE4D6")
LINK_FONT = Font(color="FF0563C1", underline="single")


# ---------------------------------------------------------------- Excel lesen
def read_rows(path):
    wb = openpyxl.load_workbook(path)
    rows = []
    for ws in wb.worksheets:
        if ws.title in DOC_SHEETS:
            hdr = None
            for r in ws.iter_rows():
                v = [c.value for c in r]
                if v[0] == "Ebene":
                    hdr = {h: i for i, h in enumerate(v) if h}
                    continue
                if not hdr or not v[hdr["Bezeichnung"]]:
                    continue
                c = r[hdr["Link / Quelle"]]
                rows.append({
                    "sheet": ws.title, "row": c.row, "col": c.column_letter, "ebene": v[0],
                    "name": str(v[hdr["Bezeichnung"]]),
                    "ref": str(v[hdr["Referenz (SR/SGS/RB)"]] or ""),
                    "text": str(c.value or "").strip(),
                    "target": c.hyperlink.target if c.hyperlink else None,
                })
        elif ws.title == GIS_SHEET:
            for r in ws.iter_rows(min_row=5):
                v = [c.value for c in r]
                if not v[2]:
                    continue
                for c in r[5:9]:
                    if c.hyperlink or (c.value and str(c.value).startswith("http")):
                        rows.append({
                            "sheet": ws.title, "row": c.row, "col": c.column_letter,
                            "ebene": v[3], "name": str(v[2]), "ref": str(v[4] or ""),
                            "label": str(c.value or ""),
                            "text": str(c.value or "") if str(c.value or "").startswith("http") else "",
                            "target": c.hyperlink.target if c.hyperlink else None,
                        })
    return rows


# ---------------------------------------------------------------- HTTP
def fetch(url):
    out = {"url": url}
    try:
        r = requests.get(url, headers=UA, timeout=40, allow_redirects=True)
        out.update(status=r.status_code, final=r.url,
                   ctype=r.headers.get("content-type", "").split(";")[0])
        body = r.content
        if out["ctype"] == "application/pdf" or body[:4] == b"%PDF":
            out["ctype"] = "application/pdf"
            try:
                from pypdf import PdfReader
                pdf = PdfReader(io.BytesIO(body))
                meta = pdf.metadata or {}
                out["title"] = str(meta.get("/Title") or "")
                out["text"] = " ".join((p.extract_text() or "") for p in pdf.pages[:3])[:6000]
                out["pages"] = len(pdf.pages)
            except Exception as e:  # noqa: BLE001
                out["text"] = f"[PDF nicht lesbar: {e}]"
        else:
            html = r.text
            m = re.search(r"<title[^>]*>(.*?)</title>", html, re.S | re.I)
            out["title"] = re.sub(r"\s+", " ", unescape(m.group(1))).strip() if m else ""
            txt = re.sub(r"<(script|style)[^>]*>.*?</\1>", " ", html, flags=re.S | re.I)
            txt = re.sub(r"<[^>]+>", " ", txt)
            out["text"] = re.sub(r"\s+", " ", unescape(txt))[:6000]
    except Exception as e:  # noqa: BLE001
        out.update(status=0, error=str(e)[:200])
    return out


# ---------------------------------------------------------------- Fedlex
FEDLEX_PREFIX = """PREFIX jolux: <http://data.legilux.public.lu/resource/ontology/jolux#>
PREFIX skos: <http://www.w3.org/2004/02/skos/core#>
"""
DEU = "<http://publications.europa.eu/resource/authority/language/DEU>"


def sparql(query):
    r = requests.post(SPARQL, data={"query": query},
                      headers={**UA, "Accept": "application/sparql-results+json"}, timeout=90)
    r.raise_for_status()
    return r.json()["results"]["bindings"]


def fedlex_by_sr(srs):
    """SR-Nummer -> [{eli, title, abbr, status}]."""
    values = " ".join(f'"{s}"' for s in srs)
    q = FEDLEX_PREFIX + f"""
SELECT DISTINCT ?sr ?ca ?title ?abbr ?status WHERE {{
  VALUES ?sr {{ {values} }}
  ?ca a jolux:ConsolidationAbstract ;
      jolux:classifiedByTaxonomyEntry ?e .
  ?e skos:notation ?n . FILTER(str(?n) = ?sr)
  OPTIONAL {{ ?ca jolux:inForceStatus ?status }}
  OPTIONAL {{ ?ca jolux:isRealizedBy ?x . ?x jolux:language {DEU} .
             OPTIONAL {{ ?x jolux:title ?title }}
             OPTIONAL {{ ?x jolux:titleShort ?abbr }} }}
}}"""
    res = {}
    for b in sparql(q):
        res.setdefault(b["sr"]["value"], []).append({
            "eli": b["ca"]["value"], "title": b.get("title", {}).get("value", ""),
            "abbr": b.get("abbr", {}).get("value", ""),
            "status": b.get("status", {}).get("value", "")})
    return res


def fedlex_by_eli(elis):
    """ELI-URI (fedlex.data.admin.ch/eli/cc/...) -> {sr, title, abbr}."""
    values = " ".join(f"<{e}>" for e in elis)
    q = FEDLEX_PREFIX + f"""
SELECT DISTINCT ?ca ?n ?title ?abbr WHERE {{
  VALUES ?ca {{ {values} }}
  OPTIONAL {{ ?ca jolux:classifiedByTaxonomyEntry/skos:notation ?n }}
  OPTIONAL {{ ?ca jolux:isRealizedBy ?x . ?x jolux:language {DEU} .
             OPTIONAL {{ ?x jolux:title ?title }}
             OPTIONAL {{ ?x jolux:titleShort ?abbr }} }}
}}"""
    res = {}
    for b in sparql(q):
        d = res.setdefault(b["ca"]["value"], {"sr": [], "title": "", "abbr": ""})
        if "n" in b and b["n"]["value"] not in d["sr"]:
            d["sr"].append(b["n"]["value"])
        d["title"] = d["title"] or b.get("title", {}).get("value", "")
        d["abbr"] = d["abbr"] or b.get("abbr", {}).get("value", "")
    return res


def fedlex_current(entries):
    """Geltender Erlass zu einer SR-Nummer: in Kraft (Status 0), sonst der jüngste nicht aufgehobene."""
    if not entries:
        return None
    in_force = [e for e in entries if e["status"].endswith("/0")]
    if in_force:
        return in_force[0]
    open_ = [e for e in entries if not e["status"].endswith("/3")]
    return max(open_, key=lambda e: e["eli"]) if open_ else None


def eli_of(url):
    m = re.search(r"fedlex\.admin\.ch/eli/(cc/[^/?#]+/[^/?#]+)", url or "")
    return f"https://fedlex.data.admin.ch/eli/{m.group(1)}" if m else None


def fedlex_url(eli):
    return eli.replace("https://fedlex.data.admin.ch/", "https://www.fedlex.admin.ch/") + "/de"


# ---------------------------------------------------------------- LexWork
def lexwork_index(base):
    """Titelverzeichnis einer LexWork-Sammlung: Nummer -> (Titel, aufgehoben)."""
    r = requests.get(base + "/api/de/texts_of_law/lightweight_index",
                     headers={**UA, "Accept": "application/json"}, timeout=60)
    r.raise_for_status()
    return {x["systematic_number"]: (x["title"], bool(x.get("abrogated")))
            for v in r.json().values() for x in v}


def lexwork_detail(base, nr):
    """Einzelabfrage (auch für aufgehobene Erlasse, die im Titelverzeichnis fehlen)."""
    r = requests.get(f"{base}/api/de/texts_of_law/{nr}", headers={**UA, "Accept": "application/json"},
                     timeout=40)
    if r.status_code != 200:
        return None
    d = r.json().get("text_of_law", {})
    dates = d.get("text_of_law_dates_str") or ""
    return d.get("title", ""), "bis:" in dates, dates


def lexwork_url(base, nr):
    return f"{base}/app/de/texts_of_law/{nr}"


# ---------------------------------------------------------------- Abgleich
def keywords(name):
    abbr = [a.strip() for a in re.findall(r"\(([A-ZÄÖÜa-zäöü0-9\-\s]{2,15})\)", name)]
    abbr = [re.sub(r"\s+(BL|SG|UR)$", "", a) for a in abbr]
    abbr += re.findall(r"\b(?=\w*[a-zäöü])[A-ZÄÖÜ][a-zäöü]?[A-ZÄÖÜ][A-Za-zäöü]{0,5}\b|\b[A-ZÄÖÜ]{3,6}\b", name)
    words = [w for w in re.findall(r"[A-Za-zÄÖÜäöüß\-]{5,}", name) if w.lower() not in STOP]
    return abbr, words[:8]


def ref_numbers(ref):
    """«SR 700.1» -> [('SR', '700.1')]; mehrere Referenzen werden alle erkannt."""
    return re.findall(r"\b(SR|SGS|sGS|RB|SRS|TR)\s*(?:-Nr\.)?\s*([0-9][0-9A-Za-z.\-]*[0-9A-Za-z])", ref or "")


def title_matches(name, title):
    """Passt ein Erlasstitel zur Bezeichnung? (Abkürzung oder Stichwort aus der Bezeichnung)."""
    t = title.lower()
    abbr, words = keywords(name)
    words = [w.strip("-") for w in words if len(w.strip("-")) >= 5]
    return any(re.search(rf"\b{re.escape(a.lower())}\b", t) for a in abbr) or \
        any(w.lower() in t for w in words)


def content_match(name, page):
    hay = f"{page.get('title', '')} {page.get('text', '')}".lower()
    abbr, words = keywords(name)
    words = [w.strip("-") for w in words if len(w.strip("-")) >= 5]
    return [w for w in words if w.lower() in hay] + [a for a in abbr if a.lower() in hay], len(words)


def is_entry_page(url):
    p = urllib.parse.urlparse(url)
    return p.path in ("", "/") and not p.query and not p.fragment


def map_url(layer, col):
    center, z = MAP_VIEWS[col]
    url = f"https://map.geo.admin.ch/#/map?lang=de&center={center}&z={z}&bgLayer=ch.swisstopo.pixelkarte-grau"
    return url + (f"&layers={layer}" if layer else "")


# ---------------------------------------------------------------- Bewertung
def assess(r, ctx):
    """Bewertet eine Zeile. Setzt r['verdict'], r['new_url'], r['new_ref'], r['reason']."""
    url = r["text"] if r["text"].startswith("http") else (r["target"] or "")
    r.update(url=url, new_url=None, new_ref=None, new_layer=None, reason="", hint="")
    pages, layer_ids = ctx["pages"], ctx["layer_ids"]
    pg = pages.get(url, {})
    status = pg.get("status")
    r["status"] = status
    r["title"] = pg.get("title") or ""
    key2, key3 = (r["sheet"], r["row"]), (r["sheet"], r["row"], r["col"])
    r["hint"] = HINWEISE.get(key2, "")

    # 1. manuell recherchierte Korrektur
    korr = KORREKTUREN.get(key3) or KORREKTUREN.get(key2)
    if korr:
        new, reason, new_ref = korr
        npg = pages.get(new, {})
        m = LEXWORK_URL.match(new)
        ok_lw = m and ctx["lw_index"].get(m.group(1), {}).get(m.group(2))
        hits, _ = content_match(r["name"], npg)
        if (ok_lw and title_matches(r["name"], ok_lw[0])) or (npg.get("status") == 200 and hits) \
                or (npg.get("status") == 200 and r["sheet"] == GIS_SHEET) \
                or (npg.get("status") == 200 and "Kartenzentrum" in reason):
            r.update(verdict="korrigiert", new_url=new, new_ref=new_ref, reason=reason)
            if ok_lw:
                r["reason"] += f" (Titel lt. Sammlung: «{ok_lw[0]}»)"
        else:
            r.update(verdict="prüfen", reason=f"Korrekturvorschlag nicht bestätigt: {new} "
                                              f"(HTTP {npg.get('status')}, Treffer {hits}) – {reason}")
        return

    # 2. Blatt 7: Layer-Links aus Layer-ID und Gemeindezentrum bilden
    if r["sheet"] == GIS_SHEET and "map.geo.admin.ch" in url:
        lid = r["ref"] if r["ref"].startswith("ch.") else LAYER_ADD.get(r["name"])
        if lid in LAYER_FIX:
            r["new_layer"] = LAYER_FIX[lid]
            lid = LAYER_FIX[lid]
        new = map_url(lid, r["col"])
        if lid and layer_ids and lid not in layer_ids:
            r.update(verdict="prüfen", reason=f"Layer-ID {lid} existiert nicht in layersConfig")
            return
        why = []
        if r["new_layer"]:
            why.append(f"Layer-ID korrigiert ({r['ref']} existiert nicht)")
        if not r["ref"].startswith("ch.") and lid:
            r["new_layer"] = lid
            why.append(f"Layer-ID ergänzt ({lid}, «{layer_ids.get(lid, '')}»)")
        if url.rstrip("/") == "https://map.geo.admin.ch":
            why.append("Hyperlink führte nur auf https://map.geo.admin.ch/ (Layer und Ausschnitt fehlten)")
        r.update(verdict="korrigiert", new_url=new, reason="; ".join(why) or "Link neu gebildet")
        if not lid:
            r["hint"] = "keine eindeutige Layer-ID; Link zeigt die Gemeinde ohne Layer"
        return

    # 3. Fedlex: SR-Nummer der Referenz
    srs = [n for p, n in ref_numbers(r["ref"]) if p == "SR"]
    if len(srs) == 1 and ctx["fed_sr"].get(srs[0]):
        cur = fedlex_current(ctx["fed_sr"][srs[0]])
        if cur:
            if not title_matches(r["name"], f"{cur['title']} ({cur['abbr']})"):
                r["hint"] = (r["hint"] + "; " if r["hint"] else "") + \
                    f"Bezeichnung und Titel zu SR {srs[0]} stimmen nicht wörtlich überein – bitte prüfen"
            correct = fedlex_url(cur["eli"])
            label = f"SR {srs[0]} «{cur['title']}»"
            if eli_of(url) == cur["eli"]:
                r.update(verdict="ok", reason=f"Fedlex-Link = {label}")
            else:
                linked = ctx["fed_eli"].get(eli_of(url) or "", {})
                old = (f"Link zeigte auf ELI {eli_of(url).split('/eli/')[1]} "
                       f"(SR {', '.join(linked.get('sr', [])) or 'keine'})") if eli_of(url) \
                    else f"Link zeigte nur auf die Einstiegsseite {url}"
                r.update(verdict="korrigiert", new_url=correct, reason=f"{old}; richtig: {label}")
            return

    # 4. LexWork-Sammlungen (kantonal / Wil)
    m = LEXWORK_URL.match(url)
    if m:
        base, nr = m.group(1), m.group(2)
        entry = ctx["lw_index"].get(base, {}).get(nr)
        refs = [(p, n) for p, n in ref_numbers(r["ref"]) if p in LEXWORK]
        if entry and (title_matches(r["name"], entry[0]) or nr in [n for _, n in refs]):
            if not title_matches(r["name"], entry[0]) and not r["hint"]:
                r["hint"] = f"verlinkte Nummer {nr} ist «{entry[0]}»; Bezeichnung lautet anders"
            canon = lexwork_url(base, nr)
            if refs and nr not in [n for _, n in refs] and not r["hint"]:
                r["hint"] = f"Referenz {r['ref']} weicht von der verlinkten Nummer {nr} ab"
            if entry[1]:
                r.update(verdict="prüfen", reason=f"Link passt zu Bezeichnung/Referenz, aber Erlass {nr} "
                                                  f"«{entry[0]}» ist aufgehoben ({entry[2] if len(entry) > 2 else ''})")
            elif canon != url:
                r.update(verdict="korrigiert", new_url=canon,
                         reason=f"Link auf Standardform der Sammlung vereinheitlicht ({nr} «{entry[0]}»)")
            else:
                r.update(verdict="ok", reason=f"{nr} «{entry[0]}»")
            return
        if entry:
            r.update(verdict="prüfen" if not r["hint"] else "hinweis",
                     reason=f"verlinkte Nummer {nr} ist «{entry[0]}»")
            return
    refs = [(p, n) for p, n in ref_numbers(r["ref"]) if p in LEXWORK]
    for p, n in refs:
        entry = ctx["lw_index"].get(LEXWORK[p], {}).get(n)
        if entry and not entry[1]:
            if not title_matches(r["name"], entry[0]):
                r["hint"] = (r["hint"] + "; " if r["hint"] else "") + \
                    f"{p} {n} ist «{entry[0]}»; Bezeichnung lautet anders – bitte prüfen"
            r.update(verdict="korrigiert", new_url=lexwork_url(LEXWORK[p], n),
                     reason=f"Link zeigte auf {url} (HTTP {status}, «{r['title'][:60]}»); "
                            f"richtig: {p} {n} «{entry[0]}»")
            return

    # 5. map.geo.admin.ch in den Blättern 1–6
    if "map.geo.admin.ch" in url:
        lids = re.findall(r"layers=([^&#]+)", url)
        fixed = url
        why = []
        for lid in lids:
            if lid in LAYER_FIX:
                fixed = fixed.replace(lid, LAYER_FIX[lid])
                r["new_layer"] = LAYER_FIX[lid]
                why.append(f"Layer-ID {lid} existiert nicht, richtig: {LAYER_FIX[lid]}")
            elif layer_ids and lid not in layer_ids:
                why.append(f"Layer-ID {lid} existiert nicht")
        if r["target"] and r["target"] != fixed:
            why.append(f"Hyperlink führte nur auf {r['target']} statt auf den angezeigten Link")
        if fixed != url or (r["target"] and r["target"] != fixed):
            r.update(verdict="korrigiert", new_url=fixed, reason="; ".join(why))
        elif lids:
            r.update(verdict="ok", reason="Layer: " + ", ".join(f"{l} «{layer_ids.get(l, '?')}»" for l in lids))
        else:
            r.update(verdict="einstieg", reason="Kartenviewer ohne Layer")
        return

    # 6. übrige Links: Erreichbarkeit und Inhalt
    hits, n = content_match(r["name"], pages.get(url, {}))
    if status == 403:
        r.update(verdict="prüfen", reason="Server blockiert automatische Abfragen (HTTP 403); im Browser prüfen")
    elif status != 200:
        r.update(verdict="fehler", reason=f"HTTP {status} {pg.get('error', '')}")
    elif is_entry_page(url):
        r.update(verdict="einstieg", reason=f"nur Einstiegsseite «{r['title'][:60]}»")
    elif hits:
        r.update(verdict="ok", reason=f"Inhalt passt (Treffer: {', '.join(hits[:4])})")
    else:
        r.update(verdict="prüfen", reason=f"erreichbar, Bezeichnung im Dokument nicht gefunden («{r['title'][:60]}»)")
    if r["hint"] and r["verdict"] == "ok":
        r["verdict"] = "hinweis"


# ---------------------------------------------------------------- Excel schreiben
VERDICT_TEXT = {"ok": "geprüft · OK", "korrigiert": "geprüft · korrigiert", "hinweis": "geprüft · Hinweis",
                "einstieg": "geprüft · nur Einstiegsseite", "prüfen": "geprüft · manuell prüfen",
                "fehler": "geprüft · Fehler"}
VERDICT_FILL = {"ok": FILL_OK, "korrigiert": FILL_KORR, "hinweis": FILL_WARN, "einstieg": None,
                "prüfen": FILL_WARN, "fehler": FILL_WARN}


def apply(path, out, rows, stamp):
    wb = openpyxl.load_workbook(path)
    for r in rows:
        ws = wb[r["sheet"]]
        cell = ws[f"{r['col']}{r['row']}"]
        if r["sheet"] == GIS_SHEET:
            if r["new_url"]:
                cell.hyperlink = r["new_url"]
                cell.font = LINK_FONT
                cell.fill = FILL_KORR
            if r["new_layer"] and r["col"] == "F":
                ws[f"E{r['row']}"].value = r["new_layer"]
                ws[f"E{r['row']}"].fill = FILL_KORR
            continue
        hdr = {c.value: c.column_letter for row in ws.iter_rows(max_row=8) for c in row if c.value}
        st, ref, note = (ws[f"{hdr[h]}{r['row']}"] for h in ("Link-Status", "Referenz (SR/SGS/RB)",
                                                             "Prüfnachweis / Hinweis"))
        if r["new_url"]:
            cell.value = r["new_url"]
            cell.hyperlink = r["new_url"]
            cell.font = LINK_FONT
            cell.fill = FILL_KORR
        if r["new_ref"] and r["new_ref"] != r["ref"]:
            ref.value = r["new_ref"]
            ref.fill = FILL_KORR
        origin = "" if r["new_url"] else str(st.value or "").split("·")[0].strip()
        st.value = f"{origin} · {VERDICT_TEXT[r['verdict']]}" if origin else VERDICT_TEXT[r["verdict"]]
        if VERDICT_FILL[r["verdict"]]:
            st.fill = VERDICT_FILL[r["verdict"]]
        msg = []
        if r["new_url"]:
            msg.append(f"Linkprüfung {stamp}: {r['reason']}. Alter Link: {r['url']}")
            if r["new_ref"] and r["new_ref"] != r["ref"] and r["ref"].strip() not in ("", "–", "—"):
                msg.append(f"Referenz alt: {r['ref']}")
        elif r["verdict"] in ("prüfen", "fehler", "hinweis"):
            msg.append(f"Linkprüfung {stamp}: {r['reason']}")
        if r["hint"]:
            msg.append(r["hint"])
        if msg:
            note.value = " | ".join(([str(note.value)] if note.value else []) + msg)

    ws = wb.create_sheet("9_Linkpruefung")
    ws.append([f"Linkprüfung {stamp} – GitHub-Runner «Linkprüfung Excel» (linkpruefung_excel.py)"])
    ws.append(["Prüfquellen: Fedlex-SPARQL (SR-Nummer ↔ Erlass), LexWork-Titelverzeichnisse (BL, SG, UR, Wil), "
               "api3.geo.admin.ch layersConfig, HTTP-Abruf mit Titel-/Textabgleich (HTML und PDF)."])
    ws.append([])
    ws.append(["Blatt", "Zeile", "Spalte", "Bezeichnung", "Referenz", "Ergebnis", "Begründung / Hinweis",
               "HTTP", "Link alt", "Link neu"])
    for c in ws[4]:
        c.font = Font(bold=True)
    order = {"korrigiert": 0, "fehler": 1, "prüfen": 2, "hinweis": 3, "einstieg": 4, "ok": 5}
    for r in sorted(rows, key=lambda x: (order[x["verdict"]], x["sheet"], x["row"], x["col"])):
        ws.append([r["sheet"], r["row"], r["col"], r["name"], r["ref"], VERDICT_TEXT[r["verdict"]],
                   "; ".join(x for x in (r["reason"], r["hint"]) if x), r["status"],
                   r["url"], r["new_url"] or ""])
        if VERDICT_FILL[r["verdict"]]:
            ws.cell(ws.max_row, 6).fill = VERDICT_FILL[r["verdict"]]
        if r["new_url"]:
            ws.cell(ws.max_row, 10).hyperlink = r["new_url"]
            ws.cell(ws.max_row, 10).font = LINK_FONT
    for col, w in zip("ABCDEFGHIJ", (20, 7, 7, 50, 18, 26, 80, 6, 50, 50)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:J{ws.max_row}"
    wb.save(out)


# ---------------------------------------------------------------- Hauptlauf
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--json", default="linkpruefung_excel.json")
    ap.add_argument("--md", default="Linkpruefung_Excel.md")
    ap.add_argument("--apply", action="store_true", help="korrigierte Excel-Datei *_geprueft.xlsx schreiben")
    args = ap.parse_args()
    stamp = datetime.date.today().strftime("%d.%m.%Y")

    rows = read_rows(args.xlsx)
    urls = {u for r in rows for u in (r["text"], r["target"]) if u and u.startswith("http")}
    urls |= {k[0] for k in KORREKTUREN.values()}
    print(f"{len(rows)} Zeilen, {len(urls)} URLs", flush=True)
    with ThreadPoolExecutor(8) as ex:
        pages = {p["url"]: p for p in ex.map(fetch, sorted(urls))}

    srs = sorted({n for r in rows for p, n in ref_numbers(r["ref"]) if p == "SR"})
    elis = sorted({e for u in urls if (e := eli_of(u))})
    fed_sr, fed_eli = fedlex_by_sr(srs), fedlex_by_eli(elis)
    lw_index = {}
    for base in LEXWORK.values():
        lw_index[base] = lexwork_index(base)
        print(f"{base}: {len(lw_index[base])} Erlasse", flush=True)
    wanted = {(m.group(1), m.group(2)) for u in urls if (m := LEXWORK_URL.match(u))}
    wanted |= {(LEXWORK[p], n) for r in rows for p, n in ref_numbers(r["ref"]) if p in LEXWORK}
    for base, nr in sorted(wanted):
        if nr not in lw_index.get(base, {}):
            det = lexwork_detail(base, nr)
            print(f"{base} {nr} nicht im Verzeichnis -> {det}", flush=True)
            if det:
                lw_index[base][nr] = det
    cfg = requests.get(LAYERS_CONFIG, headers=UA, timeout=90).json()
    layer_ids = {k: v.get("label", "") for k, v in cfg.items()}
    ctx = {"pages": pages, "fed_sr": fed_sr, "fed_eli": fed_eli, "lw_index": lw_index,
           "layer_ids": layer_ids}

    for r in rows:
        assess(r, ctx)

    with open(args.json, "w", encoding="utf-8") as f:
        json.dump({"stand": stamp, "rows": rows,
                   "pages": {u: {k: v for k, v in p.items() if k != "text"} for u, p in pages.items()},
                   "fedlex_sr": fed_sr}, f, ensure_ascii=False, indent=1)

    cnt = {}
    for r in rows:
        cnt[r["verdict"]] = cnt.get(r["verdict"], 0) + 1
    cell = lambda s: str(s or "").replace("|", "/").replace("\n", " ")  # noqa: E731
    loc = lambda r: f"{r['sheet']}!{r['col']}{r['row']}"  # noqa: E731
    lines = [f"# Linkprüfung Excel – Stand {stamp}", "",
             "Geprüft im GitHub-Runner «Linkprüfung Excel» (`.github/workflows/linkpruefung_excel.yml`, "
             "Skript `linkpruefung_excel.py`).", "",
             "Ergebnis: " + ", ".join(f"{VERDICT_TEXT[k].split('· ')[1]} {v}" for k, v in sorted(cnt.items())), "",
             "## Korrigierte Links", "", "| Zelle | Bezeichnung | Referenz | Begründung | neuer Link |",
             "| --- | --- | --- | --- | --- |"]
    for r in rows:
        if r["verdict"] == "korrigiert" and r["sheet"] != GIS_SHEET:
            ref = r["ref"] + (f" → {r['new_ref']}" if r["new_ref"] and r["new_ref"] != r["ref"] else "")
            lines.append(f"| {loc(r)} | {cell(r['name'])[:70]} | {cell(ref)} | {cell(r['reason'])} | {r['new_url']} |")
    gis = [r for r in rows if r["verdict"] == "korrigiert" and r["sheet"] == GIS_SHEET]
    lines += ["", f"## Blatt 7_GIS_Layers ({len(gis)} Links korrigiert)", "",
              "| Zelle | Layer / Datensatz | Begründung | neuer Link |", "| --- | --- | --- | --- |"]
    lines += [f"| {loc(r)} | {cell(r['name'])[:60]} | {cell(r['reason'])} | {r['new_url']} |" for r in gis]
    lines += ["", "## Manuell prüfen / Hinweise", "", "| Zelle | Bezeichnung | Ergebnis | Begründung |",
              "| --- | --- | --- | --- |"]
    lines += [f"| {loc(r)} | {cell(r['name'])[:70]} | {r['verdict']} | {cell('; '.join(x for x in (r['reason'], r['hint']) if x))} |"
              for r in rows if r["verdict"] in ("prüfen", "fehler", "hinweis")]
    lines += ["", "## Nur Einstiegsseite (erreichbar, aber kein Dokument-Link)", "",
              "| Zelle | Bezeichnung | Link |", "| --- | --- | --- |"]
    lines += [f"| {loc(r)} | {cell(r['name'])[:70]} | {r['url']} |" for r in rows if r["verdict"] == "einstieg"]
    lines += ["", "## Bestätigt", "", "| Zelle | Bezeichnung | Nachweis |", "| --- | --- | --- |"]
    lines += [f"| {loc(r)} | {cell(r['name'])[:70]} | {cell(r['reason'])} |" for r in rows if r["verdict"] == "ok"]
    with open(args.md, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print("\n".join(lines[:60]))

    if args.apply:
        out = re.sub(r"\.xlsx$", "_geprueft.xlsx", args.xlsx)
        apply(args.xlsx, out, rows, stamp)
        print("geschrieben:", out)


if __name__ == "__main__":
    sys.exit(main())
